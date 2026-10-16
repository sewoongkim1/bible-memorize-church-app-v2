# -*- coding: utf-8 -*-
"""
네비게이토 성경암송 — 엑셀 검사 → 시드 SQL(2026-10-13 · 설계 docs/superpowers/specs/2026-10-13-navigators-design.md §5)

입력: 자료/navi/네비게이토_구절입력_73.xlsx 의 시트 「구절 입력」(tools/navi-form-gen.py 가 만든 양식 · 2행부터)
       A 순번 · B 번호 · C 묶음 코드 · D 묶음 · E 주제 · F 카드 번호 · G 장절(짧게) · H 장절(풀어) · I 본문 · M 메모
출력: supabase/navi_verses.sql (--sql 일 때만) — 표 칸 둘(chapter·topic) + 카드 줄

⚠️ 본문은 사람이 붙여 넣은 원문(대한성서공회 개역개정 4판) 그대로다 — 이 도구는 **공백과 문장부호만** 다듬는다(앞뒤 · 줄바꿈 · 겹공백 · NBSP ·
   쉼표 같은 문장부호 — 2026-10-16 친구 「암송 구절에 쉼표 등은 제외」 · 주간 말씀 본문도 문장부호가 없다 · 빈칸·카드에 쉼표가 붙어 나왔다).
   낱말은 바꾸지 않는다. 숫자·괄호·영문이 들어 있으면 고치지 않고 멈춘다(붙여 넣다 따라온 절 번호·각주일 수 있다).
⚠️ 번호(no)는 진도의 열쇠다 — 이미 만든 supabase/navi_verses.sql 과 견줘, 같은 번호가 다른 장절을 가리키게 되면 멈춘다(--force 로만 강행).
   운영에 들어간 뒤 번호를 바꾸면 이미 외운 분의 기록이 엉뚱한 구절에 붙는다(되돌릴 수 없다). 카드를 내릴 때는 M열 메모에 「빼기」.
⚠️ 시드는 줄을 **켠 채**로 넣고, 다시 돌려도 is_active 를 덮지 않는다 — 문은 앱의 naviVisible(app_config naviPublic)이다(docs/notes/navigators.md).

사용법:
  python tools/navi-fit.py            # 검사만(보고서)
  python tools/navi-fit.py --sql      # 막는 것이 없으면 supabase/navi_verses.sql 을 만든다
  python tools/navi-fit.py --sql --force
"""
import argparse
import datetime
import importlib.util
import os
import re
import sys

from openpyxl import load_workbook

for _s in (sys.stdout, sys.stderr):
    if hasattr(_s, "reconfigure"):
        _s.reconfigure(encoding="utf-8")   # SystemExit 문구가 cp949 콘솔에서 깨지지 않게(psalm-fit.py 와 같다)

ROOT = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
DEF_XLSX = os.path.join(ROOT, "자료", "navi", "네비게이토_구절입력_73.xlsx")
OUT_SQL = os.path.join(ROOT, "supabase", "navi_verses.sql")
SHEET = "구절 입력"
FIRST_ROW = 2

# 양식 생성기의 묶음·책 이름·장절 풀이를 그대로 쓴다(두 곳에 따로 두면 한쪽만 고친다) — 파일 이름에 하이픈이 있어 importlib 로.
_spec = importlib.util.spec_from_file_location("navi_form_gen", os.path.join(ROOT, "tools", "navi-form-gen.py"))
_gen = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_gen)
BOOK, CHAPTERS, ref_full = _gen.BOOK, _gen.CHAPTERS, _gen.ref_full

# 묶음 코드 → 번호대(첫 번호 ~ 다음 묶음 첫 번호 - 1 · 마지막 묶음은 첫 번호 + 19)
BANDS = {}
for _i, (_code, _title, _first, _prefix, _items) in enumerate(CHAPTERS):
    _end = CHAPTERS[_i + 1][2] - 1 if _i + 1 < len(CHAPTERS) else _first + 19
    BANDS[_code] = (_first, _end)
NO_MIN, NO_MAX = min(a for a, _ in BANDS.values()), max(b for _, b in BANDS.values())

RE_SHORT = re.compile(r"^(\S+) (\d+):(\d+(?:-\d+)?(?:, ?\d+(?:-\d+)?)*)$")
RE_BAD_TEXT = re.compile(r"[0-9０-９()\[\]{}<>（）【】A-Za-z]")


RE_PUNCT = re.compile(r"[,.;:!?·…\"'“”‘’「」『』]")


def clean_text(s):
    """공백 · 문장부호만 다듬는다. 반환: (다듬은 글, [무엇을 고쳤나])"""
    fixes = []
    t = str(s or "")
    if RE_PUNCT.search(t):
        fixes.append("문장부호 " + "".join(sorted(set(RE_PUNCT.findall(t)))))
        t = RE_PUNCT.sub("", t)
    if t != t.strip():
        fixes.append("앞뒤 공백")
    t = t.strip()
    if re.search(r"[\n\r\t 　]", t):
        fixes.append("줄바꿈·탭·특수공백")
        t = re.sub(r"[\n\r\t 　]+", " ", t)
    if re.search(r" {2,}", t):
        fixes.append("겹공백")
        t = re.sub(r" {2,}", " ", t)
    return t, fixes


def read_rows(path):
    wb = load_workbook(path)   # 수식 칸(J·K)은 쓰지 않는다 — 글자 수는 여기서 센다
    if SHEET not in wb.sheetnames:
        raise SystemExit(f"「{SHEET}」 시트가 없습니다: {path}")
    ws = wb[SHEET]
    rows = []
    for r in range(FIRST_ROW, ws.max_row + 1):
        cell = lambda c: ws.cell(r, c).value
        if cell(1) is None and cell(2) is None and not cell(7):
            continue
        rows.append({
            "row": r,
            "seq": cell(1),
            "no": cell(2),
            "code": str(cell(3) or "").strip(),
            "topic": str(cell(5) or "").strip(),
            "short": str(cell(7) or "").strip(),
            "full": str(cell(8) or "").strip(),
            "text": cell(9),
            "memo": str(cell(13) or "").strip(),
        })
    return rows


def sql_escape(s):
    return str(s).replace("'", "''")


# 이미 만든 SQL 의 (번호 → 장절) — 값 줄 「  (2001, 'navi', 'assurance', '구원의 확신', '요일 5:11-12', …」 의 앞머리만 본다.
RE_NO_REF = re.compile(r"^\s*\(\s*(\d+)\s*,\s*'navi'\s*,\s*'(?:[^']|'')*'\s*,\s*'(?:[^']|'')*'\s*,\s*'((?:[^']|'')*)'", re.M)


def read_existing_no_ref(path):
    if not os.path.exists(path):
        return {}
    with open(path, encoding="utf-8") as f:
        content = f.read()
    return {int(m.group(1)): m.group(2).replace("''", "'") for m in RE_NO_REF.finditer(content)}


def build_sql(cards, today):
    on = sum(1 for c in cards if c["active"])
    lines = [
        "-- 📚 네비게이토 성경암송 — 표 칸 둘 + 카드 줄(2026-10-13 · 설계 docs/superpowers/specs/2026-10-13-navigators-design.md)",
        "-- tools/navi-fit.py 가 만든 파일이다. 손으로 고치지 말고 엑셀(자료/navi/네비게이토_구절입력_73.xlsx)을 고쳐 다시 돌린다.",
        f"-- 만든 날 {today} · 카드 {len(cards)}장(켬 {on} · 내림 {len(cards) - on}) · 번호 {cards[0]['no']}~{cards[-1]['no']}",
        "-- 본문: 대한성서공회 개역개정 4판 원문(엑셀 I열 · 사람이 붙여 넣고 원문과 대조) — 도구는 공백 · 문장부호(쉼표 등 — 2026-10-16 친구)만 다듬는다.",
        "--",
        "-- ⚠️ 개발 DB(ktpwthwqzgcqcrmsafdo)에서 먼저 돌린 뒤 운영(xnomlgydifiqiybervtf) · 넣은 뒤 python tools/sql-applied.py record supabase/navi_verses.sql --dev|--prod",
        "-- ⚠️ 차례: api(getVerses navi 갈래) → 앱 → 이 SQL. 거꾸로 돌아도 깨지지 않는다 — 주간을 읽는 곳은 모두 track='weekly' 로 거른다",
        "--    (getVerses 기본 · latestVerse · weeklyVerseKst · monitor · verseNextNo). 옛 앱은 track 을 안 보내 주간만 받는다.",
        "-- ⚠️ 줄은 켠 채(is_active=true)로 넣는다 — 문은 앱의 naviVisible(app_config naviPublic · 🧪 시험 참여자)이다(docs/notes/navigators.md).",
        "--    다시 돌려도 is_active 는 **덮지 않는다**(on conflict 갈래에 없다) — 내려 둔 카드가 되살아나지 않는다.",
        "--    엑셀 메모 「빼기」(is_active=false)는 **처음 넣을 때만** 먹는다. 이미 들어간 카드를 내리려면 손으로:",
        "--    update public.verses set is_active = false where no = <번호> and track = 'navi';",
        "-- ⚠️ 번호(no)는 진도의 열쇠다 — 한 번 넣은 번호는 다른 구절에 다시 쓰지 않는다(도구가 이 파일과 견줘 막는다).",
        "--    verses 에서 delete 금지 — progress·challenge_log·reviews 가 on delete cascade 라 성도님 기록이 함께 지워진다.",
        "-- ⚠️ 여러 번 돌려도 안전하다.",
        "",
        "begin;",
        "-- 성도님이 쓰는 중에 표 잠금을 오래 기다리지 않게(member_merge.sql 과 같다) — 5초 안에 못 잡으면 통째로 되돌리고 멈춘다.",
        "set local lock_timeout = '5s';",
        "",
        "-- ① 표 확장 -----------------------------------------------------------",
        "alter table public.verses",
        "  add column if not exists chapter text,   -- 네비게이토 묶음 코드: assurance · guide · A~E (주간·물가는 null)",
        "  add column if not exists topic   text;   -- 네비게이토 주제 이름 (주간·물가는 null)",
        "",
        "-- ② 카드 ---------------------------------------------------------------",
        "insert into public.verses (no, track, chapter, topic, ref, ref_short, ref_full, text, week, is_active)",
        "values",
    ]
    vals = []
    for c in cards:
        vals.append(
            f"  ({c['no']}, 'navi', '{sql_escape(c['code'])}', '{sql_escape(c['topic'])}', "
            f"'{sql_escape(c['short'])}', '{sql_escape(c['short'])}', '{sql_escape(c['full'])}', "
            f"'{sql_escape(c['text'])}', null, {'true' if c['active'] else 'false'})"
        )
    lines.append(",\n".join(vals))
    lines += [
        "on conflict (no) do update set",
        "  track = excluded.track, chapter = excluded.chapter, topic = excluded.topic,",
        "  ref = excluded.ref, ref_short = excluded.ref_short, ref_full = excluded.ref_full,",
        "  text = excluded.text;",
        "",
        "commit;",
        "",
        "-- ③ 확인 ---------------------------------------------------------------",
        "-- 갈래마다 줄 수 · 켠 수 · 번호대 — navi 는 " + f"{len(cards)}줄 · {cards[0]['no']}~{cards[-1]['no']} · 주간은 1000 아래 그대로",
        "select track, count(*) as 줄, count(*) filter (where is_active) as 켬, min(no) as 첫번호, max(no) as 끝번호",
        "  from public.verses group by track order by track;",
        "",
    ]
    return "\n".join(lines) + "\n"


def main():
    ap = argparse.ArgumentParser(description="네비게이토 엑셀 검사 → supabase/navi_verses.sql")
    ap.add_argument("--xlsx", default=DEF_XLSX)
    ap.add_argument("--sql", action="store_true", help="막는 것이 없으면 시드 SQL 을 만든다")
    ap.add_argument("--force", action="store_true", help="이미 만든 번호가 다른 장절을 가리키게 돼도 강행한다")
    args = ap.parse_args()

    rows = read_rows(args.xlsx)
    bar = "─" * 64
    block, warn, cards = [], [], []
    seen_no, seen_ref = {}, {}
    for r in rows:
        where = f"{r['row']}행"
        try:
            no = int(r["no"])
        except (TypeError, ValueError):
            block.append(f"{where} 번호가 숫자가 아님: {r['no']!r}")
            continue
        code = r["code"]
        if code not in BANDS:
            block.append(f"{where} no={no} 묶음 코드를 모름: {code!r} (assurance · guide · A~E)")
            continue
        a, b = BANDS[code]
        if not (a <= no <= b):
            block.append(f"{where} no={no} 이 묶음 {code} 의 번호대 {a}~{b} 밖")
        if not (NO_MIN <= no <= NO_MAX):
            block.append(f"{where} no={no} 이 네비게이토 번호대 {NO_MIN}~{NO_MAX} 밖(1~999 주간 · 1001~1999 쉴만한 물가)")
        if no in seen_no:
            block.append(f"{where} no={no} 겹침 — {seen_no[no]}행과 같은 번호")
        seen_no[no] = r["row"]
        if not r["topic"]:
            block.append(f"{where} no={no} 주제(E열)가 비었음")
        m = RE_SHORT.match(r["short"])
        if not m or m.group(1) not in BOOK:
            block.append(f"{where} no={no} 장절(짧게) 꼴이 아님: {r['short']!r}")
        else:
            if r["full"] != ref_full(r["short"]):
                block.append(f"{where} no={no} 장절(풀어) {r['full']!r} ≠ {ref_full(r['short'])!r}")
            if r["short"] in seen_ref:
                block.append(f"{where} no={no} 장절 {r['short']} 겹침 — {seen_ref[r['short']]}행")
            seen_ref[r["short"]] = r["row"]
        text, fixes = clean_text(r["text"])
        if not text:
            block.append(f"{where} no={no} 본문(I열)이 비었음")
        elif RE_BAD_TEXT.search(text):
            block.append(f"{where} no={no} 본문에 숫자·괄호·영문 — 붙여 넣다 따라온 절 번호·각주인지 원문과 대조: {RE_BAD_TEXT.findall(text)}")
        if fixes:
            warn.append(f"{where} no={no} 다듬음: {', '.join(fixes)}")
        active = "빼기" not in r["memo"]
        if not active:
            warn.append(f"{where} no={no} 메모 「빼기」 — 꺼진 채(is_active=false)로 넣는다(처음 넣을 때만)")
        cards.append({"no": no, "code": code, "topic": r["topic"], "short": r["short"], "full": r["full"],
                      "text": text, "active": active, "len": len(text.replace(" ", ""))})

    cards.sort(key=lambda c: c["no"])
    print(bar)
    print(f" 네비게이토 엑셀 검사 — {os.path.relpath(args.xlsx, ROOT)}")
    print(f" 카드 {len(cards)}장 · 막는 것 {len(block)} · 알림 {len(warn)}")
    by = {}
    for c in cards:
        by.setdefault(c["code"], []).append(c)
    for code, *_ in CHAPTERS:
        cs = by.get(code, [])
        if cs:
            print(f"   {code:9s} {len(cs):2d}장  {cs[0]['no']}~{cs[-1]['no']}")
    if cards:
        lo = min(cards, key=lambda c: c["len"])
        hi = max(cards, key=lambda c: c["len"])
        print(f" 글자 수(공백 빼고): 가장 짧은 {lo['len']}({lo['short']}) · 가장 긴 {hi['len']}({hi['short']})")
    for title, items in (("막는 것", block), ("알림", warn)):
        if items:
            print(bar)
            print(f" {title}")
            for x in items[:40]:
                print("   " + x)
            if len(items) > 40:
                print(f"   … 그 밖에 {len(items) - 40}건")
    print(bar)

    if args.sql:
        if block:
            raise SystemExit(" 고칠 곳이 남아 시드를 만들지 않았습니다.")
        if not cards:
            raise SystemExit(" 카드가 없어 시드를 만들지 않았습니다.")
        existing = read_existing_no_ref(OUT_SQL)
        changed = [(c["no"], existing[c["no"]], c["short"]) for c in cards
                   if c["no"] in existing and existing[c["no"]] != c["short"]]
        gone = sorted(set(existing) - {c["no"] for c in cards})
        if changed or gone:
            print(" ⚠️ 이미 만든 번호가 바뀌거나 빠집니다" + ("  → --force 라 강행합니다" if args.force else ""))
            for no, old, new in changed:
                print(f"   no={no}   {old}  →  {new}")
            for no in gone:
                print(f"   no={no}   {existing[no]}  →  (엑셀에서 빠짐 — 지우지 말고 메모 「빼기」로)")
            print(bar)
            if not args.force:
                raise SystemExit(
                    " 번호 재배정 감지 — 시드를 만들지 않았습니다.\n"
                    " 이 번호로 이미 외운 기록이 있으면 엉뚱한 구절에 붙습니다(되돌릴 수 없음).\n"
                    " 엑셀을 되돌리거나, 정말 의도한 것이면 --force 로 다시 돌리세요.")
        with open(OUT_SQL, "w", encoding="utf-8", newline="\n") as f:
            f.write(build_sql(cards, datetime.date.today().isoformat()))
        print(f" 시드를 만들었습니다: {os.path.relpath(OUT_SQL, ROOT)}  ({len(cards)}장)")
        print(" ⚠️ 개발 DB 에서 먼저 돌린 뒤 운영입니다 · 넣은 뒤 sql-applied.py record")


if __name__ == "__main__":
    main()
