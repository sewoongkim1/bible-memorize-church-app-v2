# -*- coding: utf-8 -*-
"""
네비게이토 성경암송 — 본문 입력 양식 생성기(2026-10-13 친구 · 설계 docs/superpowers/specs/2026-10-13-navigators-design.md)

목적: 5확신(5) · 생활 지침(8) · 주제별 60구절(A~E) — 카드 73장의 「묶음 · 주제 · 카드 번호 · 장절」을 미리 채운 엑셀을 만든다.
     친구(또는 담당자)는 **본문 칸 하나**에 개역개정 4판을 붙여 넣는다. 다 채운 표가 DB 시드(supabase/navi_verses.sql)의 원본이 된다.

⚠️ **본문은 사람이 붙여 넣는다.** 성경 본문은 한 글자도 틀리면 안 되는 자리라 AI 가 채우지 않는다(쉴만한 물가와 같은 규칙).
⚠️ 장절 출처: 60구절은 navigators.org 공식 TMS 페이지 + 한국 목록 둘이 같다. 5확신·생활 지침 8구절은 제3자 목록 둘이 같다(공식 원본과 대조 못 함).
   **카드 실물로 한 번 대조할 것** — 특히 생활 지침 8구절의 주제 이름(비워 둠)과 카드 번호 매기는 법(자료마다 다르다).
⚠️ 번호(no)는 진도의 열쇠다 — 한 번 넣으면 바꾸지도 지우지도 않는다(verses 를 지우면 기록이 cascade 로 지워진다).
   묶음마다 띄어 둔다(5확신 2001~ · 생활 지침 2011~ · A 2021~ · B 2041~ · C 2061~ · D 2081~ · E 2101~). 1001~1999 는 쉴만한 물가.

사용법: python tools/navi-form-gen.py
출력:   자료/navi/네비게이토_구절입력_73.xlsx
"""
import os
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

OUT_DIR = os.path.join(os.path.dirname(__file__), "..", "자료", "navi")
OUT_XLSX = os.path.join(OUT_DIR, "네비게이토_구절입력_73.xlsx")

NAVY, GOLD, CREAM = "1A3A6B", "C3A253", "FAF7F0"
HEAD_BG, SUB_BG, YEL_BG = "1A3A6B", "F2ECE0", "FFF6E0"
THIN = Side(style="thin", color="D8D0BC")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

BOOK = {"요일": "요한일서", "요": "요한복음", "고전": "고린도전서", "고후": "고린도후서", "잠": "잠언", "행": "사도행전", "롬": "로마서",
        "벧전": "베드로전서", "갈": "갈라디아서", "딤후": "디모데후서", "수": "여호수아", "빌": "빌립보서", "마": "마태복음", "히": "히브리서",
        "사": "이사야", "엡": "에베소서", "딛": "디도서", "계": "요한계시록", "애": "예레미야애가", "민": "민수기", "시": "시편", "눅": "누가복음",
        "막": "마가복음", "레": "레위기"}

# (묶음 코드, 묶음 이름, 첫 번호, 카드 번호 앞머리, [(주제, 장절 짧게), ...])
CHAPTERS = [
    ("assurance", "그리스도인의 확신", 2001, "확신", [
        ("구원의 확신", "요일 5:11-12"), ("기도 응답의 확신", "요 16:24"), ("승리의 확신", "고전 10:13"),
        ("사죄의 확신", "요일 1:9"), ("인도의 확신", "잠 3:5-6")]),
    ("guide", "그리스도인의 생활 지침", 2011, "지침", [
        ("", "요 15:5"), ("", "행 20:32"), ("", "롬 8:14"), ("", "고후 5:7"),
        ("", "요일 4:11"), ("", "요일 1:7"), ("", "벧전 3:15"), ("", "고후 9:7")]),
    ("A", "A. 새로운 삶", 2021, "A", [
        ("중심되신 그리스도", "고후 5:17"), ("중심되신 그리스도", "갈 2:20"), ("그리스도께 순종", "롬 12:1"), ("그리스도께 순종", "요 14:21"),
        ("말씀", "딤후 3:16"), ("말씀", "수 1:8"), ("기도", "요 15:7"), ("기도", "빌 4:6-7"),
        ("교제", "마 18:20"), ("교제", "히 10:24-25"), ("증거", "마 4:19"), ("증거", "롬 1:16")]),
    ("B", "B. 그리스도를 전파함", 2041, "B", [
        ("모든 사람이 죄를 범함", "롬 3:23"), ("모든 사람이 죄를 범함", "사 53:6"), ("죄의 형벌", "롬 6:23"), ("죄의 형벌", "히 9:27"),
        ("그리스도가 형벌을 받음", "롬 5:8"), ("그리스도가 형벌을 받음", "벧전 3:18"), ("선행으로 구원받지 못함", "엡 2:8-9"), ("선행으로 구원받지 못함", "딛 3:5"),
        ("그리스도를 모셔야 함", "요 1:12"), ("그리스도를 모셔야 함", "계 3:20"), ("구원의 확신", "요일 5:13"), ("구원의 확신", "요 5:24")]),
    ("C", "C. 하나님을 의뢰함", 2061, "C", [
        ("성령", "고전 3:16"), ("성령", "고전 2:12"), ("능력", "사 41:10"), ("능력", "빌 4:13"),
        ("성실", "애 3:22-23"), ("성실", "민 23:19"), ("평안", "사 26:3"), ("평안", "벧전 5:7"),
        ("공급", "롬 8:32"), ("공급", "빌 4:19"), ("유혹에서 도우심", "히 2:18"), ("유혹에서 도우심", "시 119:9, 11")]),
    ("D", "D. 그리스도 제자의 자격", 2081, "D", [
        ("그리스도를 첫자리에 모심", "마 6:33"), ("그리스도를 첫자리에 모심", "눅 9:23"), ("죄에서 떠남", "요일 2:15-16"), ("죄에서 떠남", "롬 12:2"),
        ("견고함", "고전 15:58"), ("견고함", "히 12:3"), ("다른 사람을 섬김", "막 10:45"), ("다른 사람을 섬김", "고후 4:5"),
        ("후히 드릴 것", "잠 3:9-10"), ("후히 드릴 것", "고후 9:6-7"), ("세계 비전", "행 1:8"), ("세계 비전", "마 28:19-20")]),
    ("E", "E. 그리스도를 닮아 감", 2101, "E", [
        ("사랑", "요 13:34-35"), ("사랑", "요일 3:18"), ("겸손", "빌 2:3-4"), ("겸손", "벧전 5:5-6"),
        ("순결", "엡 5:3"), ("순결", "벧전 2:11"), ("정직", "레 19:11"), ("정직", "행 24:16"),
        ("믿음", "히 11:6"), ("믿음", "롬 4:20-21"), ("선행", "갈 6:9-10"), ("선행", "마 5:16")]),
]


def ref_full(short):
    """「요일 5:11-12」 → 「요한일서 5장 11-12절」 · 시편은 「편」(앱의 verseRefFull 과 같은 꼴)"""
    book, rest = short.split(" ", 1)
    chap, verses = rest.split(":", 1)
    unit = "편" if book == "시" else "장"
    return f"{BOOK[book]} {chap}{unit} {verses}절"


def rows():
    out, seq = [], 0
    for code, title, first_no, prefix, items in CHAPTERS:
        for i, (topic, short) in enumerate(items):
            seq += 1
            out.append(dict(seq=seq, no=first_no + i, code=code, title=title, topic=topic, card=f"{prefix}-{i + 1}",
                            short=short, full=ref_full(short)))
    return out


def build_guide(wb):
    ws = wb.create_sheet("안내", 0)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 100

    def put(r, v, *, bold=False, size=11, color="1F2328", h=None):
        c = ws.cell(row=r, column=2, value=v)
        c.font = Font(bold=bold, size=size, color=color)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        if h:
            ws.row_dimensions[r].height = h

    put(2, "네비게이토 성경암송 — 본문 입력 양식(카드 73장)", bold=True, size=18, color=NAVY, h=34)
    put(3, "그리스도인의 확신 5 · 그리스도인의 생활 지침 8 · 주제별 60구절 A~E(각 12)", color="4A5560")
    put(5, "채워 주실 것은 「구절 입력」 시트의 I열 「본문」 하나입니다", bold=True, size=13, color=NAVY)
    put(6, "· 대한성서공회 성경 읽기(bible.bskorea.or.kr)에서 개역개정을 열어 그 절을 복사해 붙여 넣어 주세요.", h=22)
    put(7, "· 절 번호 · 괄호 · 각주 표시는 빼 주세요. 두 절짜리 카드(예: 요일 5:11-12)는 한 칸에 띄어쓰기 하나로 이어 붙입니다.", h=22)
    put(8, "· 시 119:9, 11 처럼 가운데 절을 건너뛰는 카드는 9절 다음에 11절을 띄어쓰기 하나로 이어 붙입니다.", h=22)
    put(9, "· J열 「글자 수」 · K열 「확인」은 저절로 채워집니다. 본문이 빈 줄은 노란색으로 보입니다.", h=22)
    put(11, "함께 확인해 주실 것(카드 실물이 있으면)", bold=True, size=13, color=NAVY)
    put(12, "· 생활 지침 8구절의 주제 이름(E열)이 비어 있습니다 — 카드에 적힌 이름을 넣어 주세요.", h=22)
    put(13, "· 카드 번호(F열)는 자료마다 매기는 법이 달라 「A-1 ~ A-12」처럼 절마다 붙여 두었습니다. 카드와 다르면 고쳐 주세요.", h=22)
    put(14, "· 5확신 · 생활 지침 8구절의 장절은 공식 원본과 대조하지 못했습니다. 카드와 다르면 G열을 고쳐 주세요.", h=22)
    put(16, "바꾸면 안 되는 것", bold=True, size=13, color=NAVY)
    put(17, "· B열 「번호」는 앱 진도의 열쇠입니다 — 고치지 마세요. 줄을 빼야 하면 지우지 말고 M열 메모에 「빼기」라고 적어 주세요.", h=22)
    put(19, "다 채우시면 이 파일을 그대로 돌려주세요 — 검사 도구가 빈칸 · 장절 · 겹침을 확인하고 앱에 넣습니다.", color="4A5560")


def build_input(wb):
    ws = wb.create_sheet("구절 입력", 1)
    heads = [("순번", 6), ("번호", 8), ("묶음 코드", 10), ("묶음", 22), ("주제", 22), ("카드 번호", 10),
             ("장절(짧게)", 14), ("장절(풀어)", 22), ("본문(개역개정 — 붙여 넣기)", 70), ("글자 수", 8), ("확인", 12),
             ("힌트(선택)", 18), ("메모", 18)]
    for i, (h, w) in enumerate(heads, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", fgColor=HEAD_BG if i != 9 else GOLD)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30
    data = rows()
    prev = None
    for r, d in enumerate(data, start=2):
        vals = [d["seq"], d["no"], d["code"], d["title"], d["topic"], d["card"], d["short"], d["full"], "",
                f'=IF(I{r}="","",LEN(SUBSTITUTE(I{r}," ","")))', f'=IF(I{r}="","본문 필요","")', "", ""]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = BOX
            c.alignment = Alignment(vertical="top", wrap_text=(i in (4, 5, 8, 9, 12, 13)), horizontal="center" if i in (1, 2, 3, 6, 10, 11) else "left")
            if i in (1, 2, 3, 4, 6, 7, 8, 10, 11):
                c.fill = PatternFill("solid", fgColor=CREAM)
        if d["code"] != prev:   # 묶음이 바뀌는 줄은 굵게
            for i in (3, 4):
                ws.cell(row=r, column=i).font = Font(bold=True)
        prev = d["code"]
        ws.row_dimensions[r].height = 42
    last = len(data) + 1
    ws.conditional_formatting.add(f"I2:I{last}", FormulaRule(formula=[f'LEN(TRIM(I2))=0'], fill=PatternFill("solid", fgColor=YEL_BG)))
    ws.conditional_formatting.add(f"E2:E{last}", FormulaRule(formula=[f'LEN(TRIM(E2))=0'], fill=PatternFill("solid", fgColor=YEL_BG)))
    ws.freeze_panes = "I2"
    return len(data)


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    build_guide(wb)
    n = build_input(wb)
    wb.save(OUT_XLSX)
    print(f"만듦: {os.path.normpath(OUT_XLSX)} · 카드 {n}장")


if __name__ == "__main__":
    main()
