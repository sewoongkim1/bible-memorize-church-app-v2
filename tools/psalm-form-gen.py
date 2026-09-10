# -*- coding: utf-8 -*-
"""
시편 말씀 액자 — 담당자 구절 입력 양식 생성기

목적: 시편 180구절의 「출처 + 개역개정 본문」을 담당자가 채워 주실 엑셀을 만든다.
     여기서 확정된 표가 곧 DB 시드(supabase/psalm_frames.sql)의 원본이 된다.

⚠️ 이 양식의 핵심은 **길이 판정을 엑셀 수식으로 넣은 것**이다.
   담당자가 본문을 붙여 넣는 순간 F열에 길이가, G열에 판정이 뜬다 —
   「너무 김」이 빨갛게 나오면 그 자리에서 다른 절로 바꾸시면 된다.
   측정 도구(tools/psalm-fit.py)를 돌려 결과를 되돌려 드리는 왕복 한 번이 통째로 없어진다.
   11일밖에 없는 일정에서 이 왕복 한 번이 하루다.

⚠️ 잣대는 **3단계(전체 빈칸) 기준**이다. 빈칸(.word-input)은 폭이 `글자수+1em`이라
   단어마다 1em씩 넓어진다 — 0단계(그냥 읽는 화면) 기준으로 5줄을 맞추면
   3단계에서 6~7줄이 되어 액자를 뚫는다. 그래서 재는 값이
   **글자수(공백 제외) + 단어수**이고 상한이 76이다.
   자세한 계산은 docs/superpowers/specs/2026-09-10-psalm-frame-design.html 06장.

⚠️ **본문은 사람이 붙여 넣는다.** 성경 본문은 한 글자도 틀리면 안 되는 자리라
   AI가 채우지 않는다 — 그럴듯하게 틀린 성경이 올라가면 되돌릴 방법이 없다.

사용법: python tools/psalm-form-gen.py
출력:   psalm/시편말씀액자_구절입력_180.xlsx
"""
import datetime as dt
import os
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

OUT_DIR = os.path.join(os.path.dirname(__file__), "..", "psalm")
OUT_XLSX = os.path.join(OUT_DIR, "시편말씀액자_구절입력_180.xlsx")

START_DATE = dt.date(2026, 9, 21)   # 1일차 (월) — app_config('psalm')의 start 와 같은 값
TOTAL_DAYS = 180                    # 180일차 = 2027-03-19 (금)
MAX_W = 76                          # 3단계 폭 상한 (글자수 + 단어수), 글씨 바닥 18px 기준
MIN_W = 30                          # 이보다 짧으면 액자가 허전하다
FRAME_COUNT = 12                    # 액자 그림 종류

WEEKDAY = "월화수목금토일"

# 색 — 앱과 같은 계열(군청 #1a3a6b / 금색 #c3a253 / 크림)
NAVY = "1A3A6B"
GOLD = "C3A253"
CREAM = "FAF7F0"
HEAD_BG = "1A3A6B"
SUB_BG = "F2ECE0"
RED_BG = "FCE4E4"
YEL_BG = "FFF6E0"
GRN_BG = "EDF6EE"

THIN = Side(style="thin", color="D8D0BC")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _head(ws, row, col, text, width=None):
    c = ws.cell(row=row, column=col, value=text)
    c.font = Font(bold=True, color="FFFFFF", size=11)
    c.fill = PatternFill("solid", fgColor=HEAD_BG)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = BOX
    if width:
        ws.column_dimensions[get_column_letter(col)].width = width
    return c


def build_guide(wb):
    """안내 시트 — 담당자가 먼저 보는 화면. 규칙 하나와 예시 다섯 줄이면 충분하다."""
    ws = wb.create_sheet("안내", 0)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 30

    def put(r, c, v, *, bold=False, size=11, color="1F2328", fill=None, span=None, wrap=False):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = Font(bold=bold, size=size, color=color)
        cell.alignment = Alignment(vertical="center", wrap_text=wrap)
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)
        if span:
            ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + span - 1)
        return cell

    ws.row_dimensions[2].height = 34
    put(2, 2, "시편 말씀 액자 — 구절 입력 양식", bold=True, size=18, color=NAVY, span=5)
    put(3, 2, f"{START_DATE:%Y년 %m월 %d일}({WEEKDAY[START_DATE.weekday()]}) 1일차 시작 · "
              f"하루에 한 편씩 {TOTAL_DAYS}일", size=11, color="4A5560", span=5)

    put(5, 2, "채워 주실 것은 두 칸입니다", bold=True, size=13, color=NAVY, span=5)
    put(6, 2, "「구절 입력」 시트의  D열 출처  ·  E열 개역개정 본문  — 이 둘뿐입니다.",
        size=11, span=5)
    put(7, 2, "나머지(순번·열리는 날·길이·판정)는 저절로 채워집니다.", size=11, color="4A5560", span=5)

    put(9, 2, "본문을 붙여 넣으면 그 자리에서 길이를 알려 드립니다", bold=True, size=13, color=NAVY, span=5)
    put(10, 2, "G열 판정이 빨간 「너무 김」이면 액자 다섯 줄에 안 들어갑니다. "
               "그 절은 빼시거나 더 짧은 절로 바꿔 주세요.", size=11, wrap=True, span=5)
    ws.row_dimensions[10].height = 30
    put(11, 2, "노란 「짧음」은 괜찮습니다 — 다만 액자가 좀 허전하니 앞뒤 절을 붙여 "
               "한 편으로 묶으셔도 좋습니다.", size=11, wrap=True, span=5)
    ws.row_dimensions[11].height = 30

    put(13, 2, "왜 이런 잣대인가", bold=True, size=13, color=NAVY, span=5)
    put(14, 2, "암송 3단계에서는 모든 낱말이 빈칸이 되는데, 빈칸은 글자보다 자리를 더 차지합니다. "
               "그래서 「글자수 + 낱말수」로 재고 76을 넘지 않게 합니다.", size=11, wrap=True, span=5)
    ws.row_dimensions[14].height = 30

    # 예시 표 — 실제 시편으로 잰 값(2026-09-10)
    hdr = ["구절", "공백 포함", "글자+낱말", "판정", "액자에서"]
    for i, h in enumerate(hdr):
        c = ws.cell(row=16, column=2 + i, value=h)
        c.font = Font(bold=True, size=10, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BOX
    rows = [
        ("시편 23편 1절", "26자", "27", "짧음", "액자가 허전하다", YEL_BG),
        ("시편 119편 105절", "26자", "27", "짧음", "액자가 허전하다", YEL_BG),
        ("시편 1편 1절", "61자", "62", "좋음", "가장 알맞다", GRN_BG),
        ("시편 1편 3절", "68자", "69", "좋음", "아슬하지만 들어간다", GRN_BG),
        ("(90자짜리 절)", "90자", "91", "너무 김", "다섯 줄을 넘는다", RED_BG),
    ]
    for r, (a, b, c_, d, e, fill) in enumerate(rows, start=17):
        for i, v in enumerate([a, b, c_, d, e]):
            cell = ws.cell(row=r, column=2 + i, value=v)
            cell.font = Font(size=10, bold=(i == 3))
            cell.alignment = Alignment(horizontal="center" if i else "left", vertical="center")
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.border = BOX

    put(24, 2, "출처는 이렇게 적어 주세요", bold=True, size=13, color=NAVY, span=5)
    put(25, 2, "시편 1편 1절     ·     시편 119편 105절     — 「편」과 「절」을 붙여 주시면 됩니다.",
        size=11, span=5)
    put(26, 2, "한 편에 두 절을 묶으실 때는:  시편 1편 1-2절", size=11, color="4A5560", span=5)

    put(28, 2, "액자 그림(H열)은 비워 두셔도 됩니다", bold=True, size=13, color=NAVY, span=5)
    put(29, 2, f"비워 두시면 1~{FRAME_COUNT}번 그림이 고르게 배정됩니다. "
               "「이 말씀엔 포도가 어울린다」 싶은 곳만 번호를 적어 주세요.", size=11, wrap=True, span=5)
    ws.row_dimensions[29].height = 30

    put(31, 2, "마감", bold=True, size=13, color="8A6216", span=5)
    put(32, 2, "9월 17일(목)까지 주시면 9월 21일(월) 시작에 맞출 수 있습니다. "
               "다 못 채우셔도 괜찮습니다 — 앞에서부터 30편만 있어도 한 달을 갑니다.",
        size=11, wrap=True, span=5)
    ws.row_dimensions[32].height = 32

    put(34, 2, "※ 본문은 성경 프로그램이나 사이트에서 복사해 붙여 넣어 주세요. "
               "한 글자도 틀리면 안 되는 자리라 사람 손으로 받습니다.",
        size=10, color="8A6216", wrap=True, span=5)
    ws.row_dimensions[34].height = 30


def build_input(wb):
    """구절 입력 시트 — 담당자가 실제로 채우는 표."""
    ws = wb.create_sheet("구절 입력")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:I1")
    t = ws["A1"]
    t.value = (f"시편 말씀 액자 · 구절 입력    │    "
               f"채우실 곳은 D열(출처)과 E열(본문) 두 칸입니다    │    "
               f"G열이 빨간 「너무 김」이면 다른 절로 바꿔 주세요")
    t.font = Font(bold=True, size=12, color="FFFFFF")
    t.fill = PatternFill("solid", fgColor=NAVY)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    cols = [
        ("순번", 7), ("열리는 날", 13), ("요일", 6),
        ("출처", 20), ("개역개정 본문", 72),
        ("길이", 8), ("판정", 12), ("액자", 7), ("비고", 22),
    ]
    for i, (name, w) in enumerate(cols, start=1):
        _head(ws, 3, i, name, width=w)
    ws.row_dimensions[3].height = 26

    # 담당자가 채우는 두 칸을 눈에 띄게(연한 금색 머리글)
    for i in (4, 5):
        c = ws.cell(row=3, column=i)
        c.fill = PatternFill("solid", fgColor=GOLD)
        c.font = Font(bold=True, color="1F2328", size=11)

    first = 4
    last = first + TOTAL_DAYS - 1
    for n in range(1, TOTAL_DAYS + 1):
        r = first + n - 1
        d = START_DATE + dt.timedelta(days=n - 1)

        ws.cell(row=r, column=1, value=n)
        ws.cell(row=r, column=2, value=f"{d:%Y-%m-%d}")
        ws.cell(row=r, column=3, value=WEEKDAY[d.weekday()])

        # F: 글자수(공백 제외) + 낱말수  ← 3단계 폭(em)
        ws.cell(row=r, column=6, value=(
            f'=IF($E{r}="","",'
            f'LEN(SUBSTITUTE($E{r}," ",""))'
            f'+LEN(TRIM($E{r}))-LEN(SUBSTITUTE(TRIM($E{r})," ",""))+1)'
        ))
        # G: 판정
        ws.cell(row=r, column=7, value=(
            f'=IF($F{r}="","",'
            f'IF($F{r}>{MAX_W},"너무 김",'
            f'IF($F{r}<{MIN_W},"짧음","좋음")))'
        ))

        for col in range(1, 10):
            cell = ws.cell(row=r, column=col)
            cell.border = BOX
            if col in (1, 2, 3, 6, 7, 8):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(vertical="center", wrap_text=(col == 5))
            if col in (4, 5):                       # 채우는 칸은 흰 바탕
                cell.fill = PatternFill("solid", fgColor="FFFFFF")
            elif col in (1, 2, 3):                  # 미리 채워진 칸은 회색빛
                cell.fill = PatternFill("solid", fgColor=SUB_BG)
            if col == 3 and d.weekday() == 6:       # 주일은 붉게
                cell.font = Font(color="B04A4A", bold=True, size=10)

        # 주일 줄은 열리는 날도 살짝 강조 — 주보로 알릴 때 자리를 잡기 쉽다
        if d.weekday() == 6:
            ws.cell(row=r, column=2).font = Font(color="B04A4A", size=10)

    # 판정에 색 — 담당자가 붙여 넣는 즉시 보인다
    rng = f"A{first}:I{last}"
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'$G{first}="너무 김"'], fill=PatternFill("solid", bgColor=RED_BG), stopIfTrue=False))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'$G{first}="짧음"'], fill=PatternFill("solid", bgColor=YEL_BG), stopIfTrue=False))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'$G{first}="좋음"'], fill=PatternFill("solid", bgColor=GRN_BG), stopIfTrue=False))

    # 액자 번호 1~12만
    dv = DataValidation(type="whole", operator="between", formula1=1, formula2=FRAME_COUNT,
                        allow_blank=True, showErrorMessage=True)
    dv.error = f"액자 그림은 1~{FRAME_COUNT} 사이 번호입니다. 모르시면 비워 두세요."
    dv.errorTitle = "액자 번호"
    ws.add_data_validation(dv)
    dv.add(f"H{first}:H{last}")

    ws.freeze_panes = f"A{first}"
    ws.auto_filter.ref = f"A3:I{last}"
    return ws


def build_summary(wb, sheet_name="구절 입력"):
    """한눈 요약 — 몇 편 채웠고 몇 편이 넘치는지. 담당자가 진행률을 스스로 본다."""
    ws = wb.create_sheet("진행 상황")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 44

    ws.merge_cells("B2:D2")
    t = ws["B2"]
    t.value = "진행 상황 (저절로 계산됩니다)"
    t.font = Font(bold=True, size=16, color=NAVY)
    ws.row_dimensions[2].height = 30

    q = f"'{sheet_name}'"
    rows = [
        ("채운 구절", f'=COUNTA({q}!$E$4:$E$183)', f"{TOTAL_DAYS}편 중 몇 편을 채우셨는지"),
        ("✅ 좋음", f'=COUNTIF({q}!$G$4:$G$183,"좋음")', "그대로 쓰시면 됩니다"),
        ("⚠️ 짧음", f'=COUNTIF({q}!$G$4:$G$183,"짧음")', "써도 되지만 액자가 조금 허전합니다"),
        ("❌ 너무 김", f'=COUNTIF({q}!$G$4:$G$183,"너무 김")',
         "다섯 줄을 넘습니다 — 바꾸거나 빼 주세요"),
        ("가장 긴 구절", f'=IF(COUNT({q}!$F$4:$F$183)=0,"",MAX({q}!$F$4:$F$183))',
         f"{MAX_W} 이하여야 합니다"),
        ("남은 편", f'={TOTAL_DAYS}-COUNTA({q}!$E$4:$E$183)',
         "다 못 채우셔도 됩니다 — 앞에서부터 30편이면 한 달을 갑니다"),
    ]
    for i, (label, formula, note) in enumerate(rows, start=4):
        a = ws.cell(row=i, column=2, value=label)
        a.font = Font(bold=True, size=12, color=NAVY)
        a.fill = PatternFill("solid", fgColor=SUB_BG)
        a.alignment = Alignment(vertical="center")
        a.border = BOX
        b = ws.cell(row=i, column=3, value=formula)
        b.font = Font(bold=True, size=14)
        b.alignment = Alignment(horizontal="center", vertical="center")
        b.border = BOX
        c = ws.cell(row=i, column=4, value=note)
        c.font = Font(size=10, color="4A5560")
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = BOX
        ws.row_dimensions[i].height = 26

    ws.merge_cells("B11:D11")
    n = ws["B11"]
    n.value = ("※ 「너무 김」이 하나도 없고 채운 편이 30편을 넘으면 보내 주셔도 됩니다. "
               "나머지는 뒤이어 채워도 앱이 안 깨집니다.")
    n.font = Font(size=10, color="8A6216")
    n.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[11].height = 32


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)

    build_guide(wb)
    build_input(wb)
    build_summary(wb)

    wb.active = 0
    wb.save(OUT_XLSX)

    end = START_DATE + dt.timedelta(days=TOTAL_DAYS - 1)
    print(f"만들었습니다: {os.path.normpath(OUT_XLSX)}")
    print(f"  시트 3장 — 안내 / 구절 입력({TOTAL_DAYS}줄) / 진행 상황")
    print(f"  {START_DATE:%Y-%m-%d}({WEEKDAY[START_DATE.weekday()]}) 1일차 → "
          f"{end:%Y-%m-%d}({WEEKDAY[end.weekday()]}) {TOTAL_DAYS}일차")
    print(f"  길이 잣대: 글자수+낱말수 {MIN_W}~{MAX_W}  (3단계 전체 빈칸 기준)")


if __name__ == "__main__":
    main()
